"""Normalize MarketScreener and Capital IQ exports into one comparable schema."""
from __future__ import annotations
import csv
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Tuple

from .sector_rules import resolve_sector_rule

ALIASES = {
 "name": ["name","company","company name","empresa","nombre"],
 "ticker": ["ticker","symbol","símbolo"],
 "isin": ["isin"],
 "cnae_code": ["cnae","cnae code","código cnae"],
 "sector": ["sector","sector bme","industry","industria"],
 "subsector": ["subsector","subsector bme","sub-industry","subindustry"],
 "levered_beta": ["beta","levered beta","beta apalancada"],
 "market_cap": ["market cap","market capitalization","capitalización","capitalizacion"],
 "financial_debt": ["financial debt","total debt","deuda financiera","debt"],
 "cash": ["cash","cash & equivalents","caja"],
 "enterprise_value": ["enterprise value","ev","valor empresa"],
 "revenue": ["revenue","sales","ventas","ingresos"],
 "ebitda": ["ebitda"],
 "ebit": ["ebit","operating income"],
 "per": ["per","per 2025","p/e","price earnings"],
 "ev_ebitda": ["ve/ebitda","ve/ebitda 2025","ev/ebitda"],
 "ev_revenue": ["ve/ventas","ve/ventas 2025","ev/revenue","ev/sales"],
 "price_to_book": ["p/vc","p/vc 2025","p/bv","price to book"],
 "fiscal_close": ["cierre fiscal","fiscal year end"],
 "publication_date": ["publicación 2025","publication date"],
 "coverage": ["cobertura","coverage"],
 "notes": ["observaciones","notes"],
 "source_url": ["fuente marketscreener","source url","url"],
}
SECTOR_HINTS = {
 "media_content": ("media","broadcast","publishing","entertainment","comunicación"),
 "advertising": ("advertising","marketing","publicidad"),
 "software_data": ("software","data","technology","it services"),
 "real_estate": ("real estate","reit","property","inmobili"),
 "construction": ("construction","engineering","construcción"),
 "industrial": ("industrial","manufactur"),
 "retail": ("retail",), "wholesale": ("wholesale","distribution"),
 "hospitality": ("hotel","restaurant","hospitality"),
 "transport_logistics": ("transport","logistics","airline"),
 "energy_utilities": ("energy","utilities","power"),
 "healthcare": ("health","pharma","medical"),
}

def _clean_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("_"," ").split())

def _parse_numeric(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int,float)):
        return float(value)
    text=str(value).strip().replace("€","").replace("$","").replace("x","").replace(" ","")
    multiplier=1
    if text[-1:].upper() in ("K","M","B"):
        multiplier={"K":1e3,"M":1e6,"B":1e9}[text[-1].upper()]
        text=text[:-1]
    if text.endswith("%"):
        text=text[:-1]
        multiplier/=100
    if "," in text and "." not in text:
        text=text.replace(",",".")
    else:
        text=text.replace(",","")
    try: return float(text)*multiplier
    except ValueError: return None

def _field(row: Mapping[str,Any], canonical: str):
    normalized={_clean_header(k):v for k,v in row.items()}
    for alias in ALIASES[canonical]:
        if alias in normalized: return normalized[alias]
    return None

def _archetype(cnae: Any, sector: str) -> str:
    if cnae:
        return resolve_sector_rule(cnae, None)["archetype"]
    lowered=(sector or "").lower()
    for archetype,hints in SECTOR_HINTS.items():
        if any(hint in lowered for hint in hints): return archetype
    return "general_business"

def normalize_rows(rows: Iterable[Mapping[str,Any]], provider: str,
                   as_of: str) -> Dict[str,Any]:
    records, rejected=[],[]
    for index,row in enumerate(rows, start=2):
        name=_field(row,"name")
        ticker=_field(row,"ticker")
        isin=_field(row,"isin")
        if not (name or ticker or isin):
            rejected.append({"row":index,"reason":"missing_identifier"}); continue
        numeric={key:_parse_numeric(_field(row,key)) for key in
                 ("levered_beta","market_cap","financial_debt","cash",
                  "enterprise_value","revenue","ebitda","ebit","per",
                  "ev_ebitda","ev_revenue","price_to_book")}
        cnae=_field(row,"cnae_code")
        sector=str(_field(row,"sector") or "")
        subsector=str(_field(row,"subsector") or "")
        record={"provider":provider,"as_of":as_of,"name":name,"ticker":ticker,
                "isin":isin,"cnae_code":str(cnae or ""), "sector":sector,
                "subsector":subsector,
                "archetype":_archetype(cnae," ".join((sector,subsector))), **numeric,
                "fiscal_close":_field(row,"fiscal_close"),
                "publication_date":_field(row,"publication_date"),
                "coverage":_field(row,"coverage"), "notes":_field(row,"notes"),
                "source_url":_field(row,"source_url")}
        record["quality"]={
          "beta_ready": numeric["levered_beta"] is not None and
                        numeric["market_cap"] not in (None,0) and
                        numeric["financial_debt"] is not None,
          "multiples_ready": any(numeric[key] is not None for key in
                                 ("per","ev_ebitda","ev_revenue","price_to_book")) or
                             (numeric["enterprise_value"] is not None and
                              (numeric["ebitda"] not in (None,0) or
                               numeric["revenue"] not in (None,0)))
        }
        records.append(record)
    return {"provider":provider,"as_of":as_of,"records":records,"rejected":rejected,
            "record_count":len(records),"rejected_count":len(rejected)}

def _rows_from_sheet(sheet) -> Tuple[List[Dict[str, Any]], int]:
    """Find the first identifier-bearing header row and return its records."""
    values=list(sheet.iter_rows(values_only=True))
    identifiers={_clean_header(alias) for key in ("name","ticker","isin")
                 for alias in ALIASES[key]}
    for offset,row in enumerate(values[:30],start=1):
        headers=[_clean_header(value) for value in row]
        if any(header in identifiers for header in headers):
            raw_headers=[value if value is not None else f"__blank_{column}"
                         for column,value in enumerate(row,start=1)]
            records=[]
            for record_values in values[offset:]:
                if not any(value is not None and str(value).strip() for value in record_values):
                    continue
                records.append(dict(zip(raw_headers,record_values)))
            return records,offset
    raise ValueError(f"No company identifier header found in sheet {sheet.title!r}")

def load_export(path: Path, provider: str, as_of: str, sheet_name: Optional[str]=None) -> Dict[str,Any]:
    suffix=path.suffix.lower()
    if suffix==".csv":
        with path.open(encoding="utf-8-sig",newline="") as handle:
            rows=list(csv.DictReader(handle))
    elif suffix in (".xlsx",".xlsm"):
        from openpyxl import load_workbook
        workbook=load_workbook(path,read_only=True,data_only=True)
        if sheet_name:
            sheet=workbook[sheet_name]
            rows,header_row=_rows_from_sheet(sheet)
        else:
            errors=[]
            for sheet in workbook.worksheets:
                try:
                    rows,header_row=_rows_from_sheet(sheet); break
                except ValueError as exc:
                    errors.append(str(exc))
            else:
                raise ValueError("; ".join(errors))
    else:
        raise ValueError("Supported formats: .csv, .xlsx, .xlsm")
    result=normalize_rows(rows,provider,as_of)
    result["source_file"]=path.name
    if suffix in (".xlsx",".xlsm"):
        result["source_sheet"]=sheet.title
        result["header_row"]=header_row
    return result
