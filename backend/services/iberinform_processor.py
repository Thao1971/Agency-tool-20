"""Iberinform Dataset Processor — Generates and processes company data.

Generates a realistic synthetic dataset based on INE DIRCE 2025 distributions
by CNAE and province. When real Iberinform data arrives, the same pipeline
processes CSV/XLSX files into iberinform_companies and companies_master.

Flow:
  1. Generate/Parse → iberinform_companies (raw normalized)
  2. Entity Resolution → companies_master (canonical, with cnae_primary + province_code)
  3. Recalculate Sector Intelligence V2 + Geo Intelligence
"""

import logging
import random
import string
import unicodedata
from typing import Dict, List, Optional
from database import db
from models import new_id, now_iso
from services.cnae_catalog import CNAE_DIVISIONS, CNAE_GROUPS, get_section_for_division
from services.geo_catalog import PROVINCES, get_ccaa_for_province

logger = logging.getLogger(__name__)


def _strip_accents_local(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


# Real Iberinform deliveries give province as a NAME (e.g. "ALAVA"), not the 2-digit
# code the rest of the app uses — PROVINCES is keyed by code with accented/bilingual
# labels (e.g. "Araba/Álava"). Build a best-effort reverse lookup once at import time;
# never fabricates a code when no match is found (falls back to storing the raw name).
_PROVINCE_NAME_TO_CODE: Dict[str, str] = {}
for _code, _info in PROVINCES.items():
    for _part in _info["label"].split("/"):
        _PROVINCE_NAME_TO_CODE[_strip_accents_local(_part).strip().upper()] = _code


def _resolve_province_code(name: str) -> str:
    if not name:
        return ""
    return _PROVINCE_NAME_TO_CODE.get(_strip_accents_local(name).strip().upper(), "")


# ══════════════════════════════════════════
# INE DIRCE 2025 — Real distribution of active companies by CNAE division
# Source: INE, DIRCE. ~3.31M active companies (2025)
# ══════════════════════════════════════════

CNAE_DIVISION_DISTRIBUTION = {
    "01": 0.028, "02": 0.002, "03": 0.002,  # A: Agriculture
    "05": 0.001, "06": 0.000, "07": 0.001, "08": 0.002, "09": 0.001,  # B: Extractive
    "10": 0.012, "11": 0.003, "12": 0.000, "13": 0.003, "14": 0.003,  # C: Manufacturing
    "15": 0.002, "16": 0.004, "17": 0.002, "18": 0.004, "19": 0.000,
    "20": 0.003, "21": 0.001, "22": 0.004, "23": 0.004, "24": 0.002,
    "25": 0.010, "26": 0.002, "27": 0.002, "28": 0.004, "29": 0.002,
    "30": 0.001, "31": 0.003, "32": 0.003, "33": 0.005,
    "35": 0.005,  # D: Energy
    "36": 0.001, "37": 0.001, "38": 0.003, "39": 0.001,  # E: Water
    "41": 0.030, "42": 0.008, "43": 0.065,  # F: Construction (10.3%)
    "45": 0.020, "46": 0.060, "47": 0.080,  # G: Commerce (16%)
    "49": 0.035, "50": 0.001, "51": 0.001, "52": 0.008, "53": 0.003,  # H: Transport
    "55": 0.012, "56": 0.065,  # I: Hospitality (7.7%)
    "58": 0.003, "59": 0.004, "60": 0.001, "61": 0.003, "62": 0.020, "63": 0.004,  # J: ICT
    "64": 0.008, "65": 0.002, "66": 0.005,  # K: Finance
    "68": 0.055,  # L: Real estate
    "69": 0.032, "70": 0.025, "71": 0.025, "72": 0.004, "73": 0.010,  # M: Professional
    "74": 0.012, "75": 0.003,
    "77": 0.005, "78": 0.004, "79": 0.005, "80": 0.003, "81": 0.012, "82": 0.010,  # N: Admin
    "84": 0.001,  # O: Public admin
    "85": 0.020,  # P: Education
    "86": 0.025, "87": 0.004, "88": 0.005,  # Q: Health
    "90": 0.006, "91": 0.001, "92": 0.002, "93": 0.010,  # R: Entertainment
    "94": 0.005, "95": 0.004, "96": 0.025,  # S: Other services
    "97": 0.001, "98": 0.000,  # T: Households
    "99": 0.000,  # U: Extraterritorial
}

# Province distribution (INE DIRCE 2025)
PROVINCE_DISTRIBUTION = {
    "28": 0.162, "08": 0.130, "46": 0.062, "41": 0.042, "29": 0.041,
    "03": 0.040, "30": 0.027, "07": 0.027, "50": 0.024, "48": 0.024,
    "33": 0.020, "35": 0.020, "38": 0.019, "17": 0.018, "43": 0.017,
    "11": 0.016, "15": 0.016, "36": 0.016, "18": 0.015, "20": 0.014,
    "14": 0.013, "31": 0.013, "47": 0.012, "04": 0.012, "45": 0.012,
    "12": 0.011, "21": 0.010, "39": 0.010, "06": 0.009, "23": 0.009,
    "02": 0.008, "13": 0.008, "25": 0.008, "01": 0.008, "26": 0.007,
    "37": 0.007, "24": 0.007, "10": 0.007, "09": 0.007, "27": 0.005,
    "22": 0.005, "32": 0.005, "34": 0.003, "40": 0.003, "44": 0.003,
    "49": 0.003, "05": 0.003, "42": 0.002, "16": 0.004, "19": 0.004,
    "51": 0.001, "52": 0.001,
}

# Revenue ranges by CNAE section (median EUR thousands)
REVENUE_RANGES = {
    "A": (50, 500), "B": (200, 5000), "C": (100, 3000), "D": (500, 10000),
    "E": (200, 2000), "F": (80, 1500), "G": (100, 2000), "H": (80, 1000),
    "I": (50, 800), "J": (100, 2000), "K": (200, 5000), "L": (50, 1500),
    "M": (50, 1000), "N": (50, 800), "O": (200, 3000), "P": (50, 500),
    "Q": (80, 1500), "R": (30, 500), "S": (20, 300), "T": (10, 50),
    "U": (100, 1000),
}

EMPLOYEE_RANGES = {
    "A": (1, 20), "B": (5, 100), "C": (3, 150), "D": (5, 200),
    "E": (3, 50), "F": (2, 50), "G": (1, 30), "H": (2, 40),
    "I": (1, 25), "J": (2, 80), "K": (3, 100), "L": (1, 10),
    "M": (1, 30), "N": (2, 50), "O": (10, 200), "P": (2, 40),
    "Q": (3, 80), "R": (1, 20), "S": (1, 10), "T": (1, 3),
    "U": (2, 30),
}

# Spanish company name components
LEGAL_FORMS = ["SL", "SA", "SLU", "SAU", "SC", "SLP", "COOP"]
NAME_PREFIXES = [
    "Grupo", "Servicios", "Soluciones", "Tecnologias", "Industrias", "Comercial",
    "Distribuciones", "Construcciones", "Inversiones", "Consultoria", "Gestion",
    "Desarrollo", "Proyectos", "Ingenieria", "Comunicaciones", "Digital",
    "Logistica", "Energia", "Alimentacion", "Transporte", "Inmobiliaria",
]


def _generate_cif() -> str:
    """Generate a realistic Spanish CIF."""
    letter = random.choice("ABCDEFGHJNPQRSUVW")
    digits = "".join([str(random.randint(0, 9)) for _ in range(7)])
    check = random.choice(string.digits + "ABCDEFGHIJ")
    return f"{letter}{digits}{check}"


def _generate_company_name(cnae_div: str) -> str:
    """Generate a realistic company name."""
    prefix = random.choice(NAME_PREFIXES)
    suffix = "".join(random.choices(string.ascii_uppercase, k=random.randint(3, 6)))
    form = random.choice(LEGAL_FORMS)
    return f"{prefix} {suffix} {form}"


def _weighted_choice(distribution: Dict[str, float]) -> str:
    """Select based on weight distribution."""
    items = list(distribution.items())
    weights = [w for _, w in items]
    total = sum(weights)
    weights = [w / total for w in weights]
    return random.choices([k for k, _ in items], weights=weights, k=1)[0]


def _pick_group_for_division(div_code: str) -> Optional[str]:
    """Pick a random CNAE group (4-digit) within a division."""
    groups = [g for g, info in CNAE_GROUPS.items() if info["division"] == div_code]
    if groups:
        return random.choice(groups)
    return None


async def generate_synthetic_dataset(count: int = 5000) -> Dict:
    """Generate a realistic synthetic Iberinform dataset and persist it."""
    now = now_iso()
    companies = []
    fiscal_years = []

    for i in range(count):
        # Pick CNAE division weighted by DIRCE distribution
        cnae_div = _weighted_choice(CNAE_DIVISION_DISTRIBUTION)
        cnae_group = _pick_group_for_division(cnae_div)
        section = get_section_for_division(cnae_div)

        # Pick province weighted by distribution
        province_code = _weighted_choice(PROVINCE_DISTRIBUTION)
        ccaa_code = get_ccaa_for_province(province_code)
        province_label = PROVINCES.get(province_code, {}).get("label", "")

        # Generate company data
        cif = _generate_cif()
        name = _generate_company_name(cnae_div)

        # Revenue and employees (realistic ranges)
        rev_range = REVENUE_RANGES.get(section, (50, 500))
        emp_range = EMPLOYEE_RANGES.get(section, (1, 20))

        # Generate 2-3 fiscal years
        years_count = random.choice([2, 2, 3])
        base_revenue = random.uniform(rev_range[0], rev_range[1]) * 1000  # EUR
        base_employees = random.randint(emp_range[0], emp_range[1])
        base_ebitda_margin = random.uniform(0.03, 0.25)

        company = {
            "company_id": new_id(),
            "cif": cif,
            "cif_normalized": cif.upper().replace("-", "").replace(" ", ""),
            "legal_name": name,
            "trade_name": name.split(" ")[1] if len(name.split(" ")) > 1 else name,
            "cnae_code": cnae_group or f"{cnae_div}00",
            "cnae_division": cnae_div,
            "cnae_section": section,
            "cnae_label": CNAE_DIVISIONS.get(cnae_div, {}).get("label", ""),
            "province_code": province_code,
            "province_name": province_label,
            "ccaa_code": ccaa_code,
            "legal_form": random.choice(LEGAL_FORMS),
            "status": "active",
            "employees_latest": base_employees,
            "revenue_latest": round(base_revenue, 2),
            "source": "iberinform_synthetic",
            "source_version": "v1.0",
            "imported_at": now,
            "updated_at": now,
        }
        companies.append(company)

        # Fiscal years
        for y_offset in range(years_count):
            year = 2024 - y_offset
            growth = 1 + random.uniform(-0.15, 0.20)
            revenue = round(base_revenue * (growth ** y_offset), 2)
            employees = max(1, base_employees + random.randint(-3, 5) * (years_count - y_offset))
            ebitda = round(revenue * base_ebitda_margin * random.uniform(0.7, 1.3), 2)

            fy = {
                "fiscal_year_id": new_id(),
                "company_id": company["company_id"],
                "cif": cif,
                "year": year,
                "revenue": revenue,
                "ebitda": ebitda,
                "ebitda_margin": round(ebitda / revenue, 4) if revenue > 0 else 0,
                "employees": employees,
                "total_assets": round(revenue * random.uniform(0.5, 2.0), 2),
                "equity": round(revenue * random.uniform(0.1, 0.8), 2),
                "net_income": round(ebitda * random.uniform(0.4, 0.9), 2),
                "source": "iberinform_synthetic",
                "imported_at": now,
            }
            fiscal_years.append(fy)

    # ── Persist to iberinform_companies ──
    if companies:
        await db.iberinform_companies.delete_many({"source": "iberinform_synthetic"})
        await db.iberinform_companies.insert_many(companies)

    # ── Persist fiscal years ──
    if fiscal_years:
        await db.iberinform_financials.delete_many({"source": "iberinform_synthetic"})
        await db.iberinform_financials.insert_many(fiscal_years)

    # ── Update companies_master with CNAE and province ──
    master_updates = await _update_companies_master(companies)

    # ── Create indexes ──
    await db.iberinform_companies.create_index("company_id", unique=True)
    await db.iberinform_companies.create_index("cif")
    await db.iberinform_companies.create_index("cnae_division")
    await db.iberinform_companies.create_index("cnae_section")
    await db.iberinform_companies.create_index("province_code")
    await db.iberinform_companies.create_index("ccaa_code")
    await db.iberinform_financials.create_index("fiscal_year_id", unique=True)
    await db.iberinform_financials.create_index("company_id")
    await db.iberinform_financials.create_index("cif")
    await db.iberinform_financials.create_index("year")

    # ── Stats ──
    cnae_coverage = set(c["cnae_division"] for c in companies)
    province_coverage = set(c["province_code"] for c in companies)
    years_covered = set(fy["year"] for fy in fiscal_years)

    return {
        "status": "completed",
        "companies_imported": len(companies),
        "fiscal_years_imported": len(fiscal_years),
        "years_covered": sorted(years_covered, reverse=True),
        "cnae_divisions_covered": len(cnae_coverage),
        "provinces_covered": len(province_coverage),
        "companies_master_updated": master_updates,
        "source": "iberinform_synthetic",
        "generated_at": now,
    }


async def _update_companies_master(companies: List[Dict]) -> int:
    """Update companies_master with CNAE and province from Iberinform data."""
    updated = 0

    for comp in companies:
        # Upsert into companies_master
        result = await db.companies_master.update_one(
            {"cif_normalized": comp["cif_normalized"]},
            {"$set": {
                "legal_name": comp["legal_name"],
                "cnae_primary": comp["cnae_division"],
                "cnae_code": comp["cnae_code"],
                "cnae_section": comp["cnae_section"],
                "cnae_label": comp["cnae_label"],
                "province_code": comp["province_code"],
                "province_name": comp["province_name"],
                "ccaa_code": comp["ccaa_code"],
                "legal_form": comp["legal_form"],
                "status": comp["status"],
                "employees_latest": comp["employees_latest"],
                "revenue_latest": comp["revenue_latest"],
                # Was hardcoded "iberinform" regardless of real vs synthetic — the
                # exact same masking bug found and fixed elsewhere in the data
                # layer (master_builder.py _build_sources). Now reflects reality.
                "data_source": comp.get("source", "iberinform"),
                "updated_at": comp["updated_at"],
            },
            "$setOnInsert": {
                "master_company_id": comp["company_id"],
                "cif": comp["cif"],
                "cif_normalized": comp["cif_normalized"],
                "normalized_name": comp["legal_name"].upper(),
                "merge_status": "auto_created",
                "confidence_score": 0.85,
                "created_at": comp["imported_at"],
            }},
            upsert=True,
        )
        if result.upserted_id or result.modified_count > 0:
            updated += 1

    return updated


async def process_real_iberinform_tab_directory(directory: str = None, source_version: str = "real-2026-07", delivery=None) -> Dict:
    """Process a real Iberinform delivery (Datos_GENERALES.tab + Datos_BALANCES.tab)
    into the LEGACY iberinform_companies / iberinform_financials / companies_master
    collections — the ones still read by Sector/Geo/Cross Intelligence, DocStudio,
    Valuo integration, and the CNMV/BME company-matching connectors.

    This is a parallel, differently-shaped delivery from Daniel's 25,000-company real
    sample (2026-07): tab-separated, English field names, financials as a long
    (BALANCE_SHEET_ITEM code, value) table rather than the single flat CSV row this
    module's older `_parse_row`/`process_real_iberinform_file` expects. Verified
    against `Diccionario_Datos_Financial_Info.pdf` + `Financial Info Translated
    V2.xlsx`: the BALANCE_SHEET_ITEM codes are the SAME numbering as
    `services.data_layer.ingestion.account_map.ACCOUNT_MAP` (10000=total_assets,
    20000=equity, 40100=revenue, 40800=depreciation, 49100=operating_income,
    49500=net_income) — reused as-is here rather than re-deriving.

    See services/data_layer/ingestion/iberinform_tab_ingest.py for the modern-schema
    (norm_company/master_companies) counterpart of this same source data — the app
    has two parallel company data models today, this function only fills the legacy
    one. Upserts by CIF (does not wipe existing data), so re-running with a bigger
    or updated delivery is safe.

    `delivery` (ZipDelivery de R2): si se pasa, GENERALES/BALANCES se leen por STREAMING
    desde el zip en R2 (sin disco) en vez de desde `directory`. Ver r2_delivery.py.
    """
    import csv
    import os
    from pymongo import UpdateOne
    from services.data_layer.ingestion.account_map import parse_amount, derive_metrics, ACCOUNT_MAP

    generales_path = os.path.join(directory, "Datos_GENERALES.tab") if directory else None
    balances_path = os.path.join(directory, "Datos_BALANCES.tab") if directory else None
    has_generales = delivery.has("Datos_GENERALES.tab") if delivery is not None else bool(generales_path and os.path.isfile(generales_path))
    has_balances = delivery.has("Datos_BALANCES.tab") if delivery is not None else bool(balances_path and os.path.isfile(balances_path))
    if not has_generales:
        return {"status": "error", "message": f"No se encontro Datos_GENERALES.tab ({generales_path or delivery.key})"}

    def _open_generales():
        if delivery is not None:
            return delivery.open_text("Datos_GENERALES.tab", encoding="latin-1")
        return open(generales_path, encoding="latin-1", errors="replace", newline="")

    def _open_balances():
        if delivery is not None:
            return delivery.open_text("Datos_BALANCES.tab", encoding="utf-8")
        return open(balances_path, encoding="utf-8", errors="replace", newline="")

    now = now_iso()
    account_codes = set(ACCOUNT_MAP.keys())

    # ── Financials pivot: (cif -> year -> {code: value}), filtered to the 6 codes
    # derive_metrics() actually reads (same reasoning as the modern ingestor: the
    # file carries ~900 possible line items per company-year, no need to keep them
    # all in memory for a 25k-company batch). ──
    fin_by_cif: Dict[str, Dict[int, Dict[str, float]]] = {}
    if has_balances:
        with _open_balances() as fh:
            reader = csv.DictReader(fh, delimiter="\t")
            for row in reader:
                cif = (row.get("REG_NUMBER") or "").strip()
                code = (row.get("BALANCE_SHEET_ITEM") or "").strip()
                if not cif or code not in account_codes:
                    continue
                year_f = parse_amount(row.get("BALANCE_SHEET_YEAR"))
                val = parse_amount(row.get("BALANCE_SHEET_ITEM_VALUE"))
                if year_f is None or val is None:
                    continue
                fin_by_cif.setdefault(cif, {}).setdefault(int(year_f), {})[code] = val
    else:
        logger.warning("process_real_iberinform_tab_directory: Datos_BALANCES.tab not found — companies will have no financials")

    # ── Company file (GENERALES) ──
    companies: List[Dict] = []
    fiscal_years: List[Dict] = []
    with _open_generales() as fh:
        reader = csv.DictReader(fh, delimiter="\t")
        for row in reader:
            cif = (row.get("REG_NUMBER") or "").strip()
            if not cif or len(cif) < 5:
                continue
            name = (row.get("COMPANY_NAME") or "").strip()
            cnae = (row.get("ACTIVITY_CODE") or "").strip()
            cnae_div = cnae[:2] if len(cnae) >= 2 else ""
            section = get_section_for_division(cnae_div) if cnae_div else None
            province_name = (row.get("PROVINCE") or "").strip()
            province_code = _resolve_province_code(province_name)

            employees = 0
            try:
                employees = int(float(row.get("EMPLOYEES") or 0))
            except (ValueError, TypeError):
                pass

            company_years = fin_by_cif.get(cif, {})
            latest_year = max(company_years.keys()) if company_years else None
            latest_metrics = derive_metrics(company_years[latest_year]) if latest_year is not None else {}

            company = {
                "company_id": new_id(),
                "cif": cif,
                "cif_normalized": cif.upper().replace("-", "").replace(" ", ""),
                "legal_name": name,
                "trade_name": (row.get("TRADE_NAME") or "").strip() or name,
                "cnae_code": cnae,
                "cnae_division": cnae_div,
                "cnae_section": section,
                "cnae_label": CNAE_DIVISIONS.get(cnae_div, {}).get("label", ""),
                "province_code": province_code,
                "province_name": PROVINCES.get(province_code, {}).get("label", "") if province_code else province_name,
                "ccaa_code": PROVINCES.get(province_code, {}).get("ccaa") if province_code else None,
                "legal_form": (row.get("SHORT_ES") or "").strip(),
                "status": "active",
                "employees_latest": employees,
                "revenue_latest": latest_metrics.get("revenue") or 0,
                "source": "iberinform",
                "source_version": source_version,
                "imported_at": now,
                "updated_at": now,
            }
            companies.append(company)

    if not companies:
        return {"status": "error", "message": "No se encontraron empresas validas en Datos_GENERALES.tab"}

    # ── Stability fix (2026-07-23): company_id used to be regenerated with new_id()
    # on EVERY delivery for EVERY company, then blindly overwritten via a plain $set
    # upsert below — so company_id was never actually stable across monthly
    # deliveries, unlike companies_master.master_company_id (already correctly
    # protected via $setOnInsert, see _update_companies_master above). Nothing in
    # this codebase currently joins on iberinform_companies.company_id as a foreign
    # key (confirmed by search), so this was a silent landmine rather than an active
    # bug — but any future code naturally assuming a field literally called
    # "company_id" is a stable identifier would get quietly wrong results after the
    # second delivery. Fixed by resolving already-existing company_id values first
    # (one bulk query, only for the CIFs in this delivery) and reusing them, so a
    # pre-existing company keeps the SAME company_id forever; only genuinely new
    # CIFs get the freshly generated one. fiscal_years (built below, AFTER this
    # resolution) then always references the true, stable company_id too.
    cif_normalized_list = [c["cif_normalized"] for c in companies]
    existing_ids: Dict[str, str] = {}
    for i in range(0, len(cif_normalized_list), 2000):
        batch = cif_normalized_list[i:i + 2000]
        async for doc in db.iberinform_companies.find(
                {"cif_normalized": {"$in": batch}}, {"_id": 0, "cif_normalized": 1, "company_id": 1}):
            if doc.get("company_id"):
                existing_ids[doc["cif_normalized"]] = doc["company_id"]
    for c in companies:
        if c["cif_normalized"] in existing_ids:
            c["company_id"] = existing_ids[c["cif_normalized"]]

    # ── Fiscal years, built AFTER company_id resolution so they always reference
    # the true stable id (not a since-discarded freshly-generated one). ──
    for c in companies:
        for year, accounts in fin_by_cif.get(c["cif"], {}).items():
            metrics = derive_metrics(accounts)
            fiscal_years.append({
                "fiscal_year_id": new_id(),
                "company_id": c["company_id"],
                "cif": c["cif"],
                "year": year,
                "revenue": metrics.get("revenue"),
                "ebitda": metrics.get("ebitda"),
                "ebitda_margin": metrics.get("ebitda_margin"),
                "employees": c["employees_latest"],
                "total_assets": metrics.get("total_assets"),
                "equity": metrics.get("equity"),
                "net_income": metrics.get("net_income"),
                "source": "iberinform",
                "imported_at": now,
            })

    # Persist via chunked bulk_write (25k+ individual round trips would be slow
    # against a real remote Mongo) — upsert by cif_normalized/cif+year+source, never
    # delete_many first, so a re-run with an updated delivery only ever adds/refreshes.
    # company_id also goes into $setOnInsert (belt-and-suspenders on top of the
    # pre-resolution above) so a genuinely new company's id, once assigned, is never
    # touched by a future $set either.
    CHUNK = 2000
    company_ops = []
    for c in companies:
        cset = {k: v for k, v in c.items() if k != "company_id"}
        company_ops.append(UpdateOne(
            {"cif_normalized": c["cif_normalized"]},
            {"$set": cset, "$setOnInsert": {"company_id": c["company_id"]}},
            upsert=True,
        ))
    for i in range(0, len(company_ops), CHUNK):
        await db.iberinform_companies.bulk_write(company_ops[i:i + CHUNK], ordered=False)

    fy_ops = [UpdateOne({"cif": fy["cif"], "year": fy["year"], "source": "iberinform"}, {"$set": fy}, upsert=True)
              for fy in fiscal_years]
    for i in range(0, len(fy_ops), CHUNK):
        await db.iberinform_financials.bulk_write(fy_ops[i:i + CHUNK], ordered=False)

    master_updates = await _update_companies_master(companies)

    return {
        "status": "completed",
        "companies_imported": len(companies),
        "fiscal_years_imported": len(fiscal_years),
        "companies_master_updated": master_updates,
        "source": "iberinform",
        "source_version": source_version,
        "generated_at": now,
    }


async def purge_synthetic_dataset() -> Dict:
    """Remove the synthetic Iberinform dataset (source == "iberinform_synthetic")
    from iberinform_companies, iberinform_financials, AND the companies_master
    records that only exist because of it.

    Intentionally conservative: only deletes companies_master docs whose
    master_company_id came from a synthetic company_id AND whose data_source is
    still "iberinform_synthetic" (i.e. nothing else has legitimately claimed/updated
    that record since). A company_id that a real delivery has since re-upserted over
    (same CIF, now source="iberinform") is left alone — real data always wins,
    nothing here can delete real data.
    """
    synthetic_ids = await db.iberinform_companies.distinct("company_id", {"source": "iberinform_synthetic"})

    companies_deleted = (await db.iberinform_companies.delete_many({"source": "iberinform_synthetic"})).deleted_count
    financials_deleted = (await db.iberinform_financials.delete_many({"source": "iberinform_synthetic"})).deleted_count

    master_deleted = 0
    if synthetic_ids:
        result = await db.companies_master.delete_many({
            "master_company_id": {"$in": synthetic_ids},
            "data_source": "iberinform_synthetic",
        })
        master_deleted = result.deleted_count

    return {
        "status": "completed",
        "iberinform_companies_deleted": companies_deleted,
        "iberinform_financials_deleted": financials_deleted,
        "companies_master_deleted": master_deleted,
    }


async def process_real_iberinform_file(file_id: str) -> Dict:
    """Process a real Iberinform CSV/XLSX file from provider_file_contents.
    
    Expected columns: CIF, NOMBRE, CNAE, PROVINCIA, INGRESOS, EMPLEADOS, EBITDA, ACTIVO, PATRIMONIO
    """
    import io

    doc = await db.provider_file_contents.find_one({"file_id": file_id})
    if not doc or not doc.get("content"):
        return {"status": "error", "message": f"File {file_id} not found or empty"}

    file_meta = await db.provider_files.find_one({"file_id": file_id}, {"_id": 0})
    filename = file_meta.get("filename", "") if file_meta else ""
    content = doc["content"]

    companies = []
    now = now_iso()

    if filename.endswith(".xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content))
            ws = wb.active
            headers = [str(ws.cell(1, c).value or "").strip().upper() for c in range(1, ws.max_column + 1)]
            for r in range(2, ws.max_row + 1):
                row = {headers[c]: ws.cell(r, c + 1).value for c in range(len(headers))}
                comp = _parse_row(row, now)
                if comp:
                    companies.append(comp)
        except Exception as e:
            return {"status": "error", "message": f"XLSX parse error: {e}"}

    elif filename.endswith(".csv"):
        import csv
        try:
            text = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text), delimiter=";")
            for row in reader:
                row_upper = {k.strip().upper(): v for k, v in row.items()}
                comp = _parse_row(row_upper, now)
                if comp:
                    companies.append(comp)
        except Exception as e:
            return {"status": "error", "message": f"CSV parse error: {e}"}
    else:
        return {"status": "error", "message": f"Unsupported format: {filename}"}

    if not companies:
        return {"status": "error", "message": "No valid company records found"}

    # Persist
    await db.iberinform_companies.insert_many(companies)
    master_updates = await _update_companies_master(companies)

    # Mark file as processed
    await db.provider_files.update_one(
        {"file_id": file_id},
        {"$set": {"status": "processed", "processed": True, "processed_at": now,
                  "records_detected": len(companies)}}
    )

    return {
        "status": "completed",
        "file_id": file_id,
        "companies_imported": len(companies),
        "companies_master_updated": master_updates,
        "processed_at": now,
    }


def _parse_row(row: Dict, now: str) -> Optional[Dict]:
    """Parse a single row from Iberinform CSV/XLSX into a company document."""
    cif = str(row.get("CIF", "") or "").strip()
    if not cif or len(cif) < 5:
        return None

    name = str(row.get("NOMBRE", row.get("RAZON_SOCIAL", row.get("DENOMINACION", ""))) or "").strip()
    cnae = str(row.get("CNAE", row.get("CNAE_2009", "")) or "").strip()
    provincia = str(row.get("PROVINCIA", row.get("PROV", "")) or "").strip()

    # Revenue
    revenue = 0
    for key in ["INGRESOS", "CIFRA_NEGOCIOS", "REVENUE", "VENTAS"]:
        if key in row and row[key]:
            try:
                revenue = float(str(row[key]).replace(",", ".").replace(" ", ""))
                break
            except (ValueError, TypeError):
                pass

    employees = 0
    for key in ["EMPLEADOS", "EMPLOYEES", "PLANTILLA"]:
        if key in row and row[key]:
            try:
                employees = int(float(str(row[key]).replace(",", ".")))
                break
            except (ValueError, TypeError):
                pass

    cnae_div = cnae[:2] if len(cnae) >= 2 else ""
    section = get_section_for_division(cnae_div) if cnae_div else None

    return {
        "company_id": new_id(),
        "cif": cif,
        "cif_normalized": cif.upper().replace("-", "").replace(" ", ""),
        "legal_name": name,
        "cnae_code": cnae,
        "cnae_division": cnae_div,
        "cnae_section": section,
        "cnae_label": CNAE_DIVISIONS.get(cnae_div, {}).get("label", ""),
        "province_code": provincia if len(provincia) == 2 else "",
        "province_name": PROVINCES.get(provincia, {}).get("label", "") if len(provincia) == 2 else provincia,
        "ccaa_code": get_ccaa_for_province(provincia) if len(provincia) == 2 else None,
        "status": "active",
        "employees_latest": employees,
        "revenue_latest": revenue,
        "source": "iberinform",
        "source_version": "v1.0",
        "imported_at": now,
        "updated_at": now,
    }
