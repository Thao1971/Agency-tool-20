"""PPTX Renderer — python-pptx based PowerPoint generation."""

import io
import logging
from typing import Dict, List, Optional
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

logger = logging.getLogger(__name__)


def render_pptx(manifest: Dict, brand: Optional[Dict] = None) -> bytes:
    """Generate a PPTX from manifest data."""
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    data = manifest.get("data_payload", {})
    blocks = manifest.get("generated_blocks", {})
    overrides = manifest.get("manual_overrides", {})

    merged = {**data}
    for k, v in blocks.items():
        merged[k] = v.get("content", v) if isinstance(v, dict) else v
    for k, v in overrides.items():
        merged[k] = v

    primary = brand.get("primary_color", "#1a1a2e") if brand else "#1a1a2e"
    secondary = brand.get("secondary_color", "#3b82f6") if brand else "#3b82f6"
    font_heading = brand.get("font_heading", "Calibri") if brand else "Calibri"
    font_body = brand.get("font_body", "Calibri") if brand else "Calibri"

    primary_rgb = _hex_to_rgb(primary)
    secondary_rgb = _hex_to_rgb(secondary)

    # Cover slide
    _add_cover_slide(prs, merged, primary_rgb, font_heading, font_body)

    # Content slides
    for section_key, title in [
        ("executive_summary", "Executive Summary"),
        ("company_overview", "Company Overview"),
        ("top_strengths", "Key Strengths"),
        ("top_risks", "Key Risks"),
        ("financial_highlights", "Financial Highlights"),
        ("market_position", "Market Position"),
        ("next_steps", "Next Steps"),
    ]:
        content = merged.get(section_key)
        if content:
            _add_content_slide(prs, title, content, primary_rgb, secondary_rgb, font_heading, font_body)

    # Legal slide
    legal = manifest.get("legal_footer")
    if legal:
        _add_legal_slide(prs, legal, font_body)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _hex_to_rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _add_cover_slide(prs, data, primary_rgb, font_heading, font_body):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Background
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = primary_rgb

    # Title
    title = data.get("title", "Report")
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.bold = True
    p.font.name = font_heading
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    subtitle = data.get("subtitle", "")
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(16)
        p2.font.color.rgb = RGBColor(200, 200, 220)
        p2.font.name = font_body
        p2.alignment = PP_ALIGN.CENTER


def _add_content_slide(prs, title, content, primary_rgb, secondary_rgb, font_heading, font_body):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title bar
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.color.rgb = primary_rgb
    p.font.bold = True
    p.font.name = font_heading

    # Content
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True

    if isinstance(content, list):
        for item in content:
            p = tf.add_paragraph()
            p.text = f"  {item}"
            p.font.size = Pt(14)
            p.font.name = font_body
            p.space_after = Pt(6)
    else:
        text = content if isinstance(content, str) else str(content)
        # Split into paragraphs
        for para_text in text.split("\n"):
            if para_text.strip():
                p = tf.add_paragraph()
                p.text = para_text.strip()
                p.font.size = Pt(13)
                p.font.name = font_body
                p.space_after = Pt(4)


def _add_legal_slide(prs, legal_text, font_body):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(11), Inches(2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = legal_text
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(160, 160, 160)
    p.font.name = font_body
    p.alignment = PP_ALIGN.CENTER
